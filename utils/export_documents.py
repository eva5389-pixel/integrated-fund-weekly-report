from __future__ import annotations

from io import BytesIO
from pathlib import Path
from uuid import uuid4
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ACCENT = "0F766E"
HEADER_FILL = "DDF5EF"
PDF_FONT = "NotoSansTC"
PDF_FONT_PATH = Path(__file__).resolve().parents[1] / "assets" / "fonts" / "NotoSansTC-Regular.ttf"


def _embed_docx_font(docx_bytes: bytes) -> bytes:
    """Embed the bundled OFL font so Chinese text remains portable."""
    font_key = uuid4()
    key_bytes = font_key.bytes_le
    font_data = bytearray(PDF_FONT_PATH.read_bytes())
    for index in range(min(32, len(font_data))):
        font_data[index] ^= key_bytes[index % 16]

    source = BytesIO(docx_bytes)
    target = BytesIO()
    relations_name = "word/_rels/fontTable.xml.rels"
    with ZipFile(source, "r") as zin, ZipFile(target, "w", ZIP_DEFLATED) as zout:
        source_names = set(zin.namelist())
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "word/fontTable.xml":
                text = data.decode("utf-8")
                font = (
                    f'<w:font w:name="{PDF_FONT}"><w:family w:val="swiss"/>'
                    f'<w:embedRegular r:id="rIdNotoSansTC" '
                    f'w:fontKey="{{{str(font_key).upper()}}}"/></w:font>'
                )
                text = text.replace("</w:fonts>", font + "</w:fonts>")
                data = text.encode("utf-8")
            elif info.filename == relations_name:
                text = data.decode("utf-8")
                relation = (
                    '<Relationship Id="rIdNotoSansTC" '
                    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/font" '
                    'Target="fonts/NotoSansTC.odttf"/>'
                )
                text = text.replace("</Relationships>", relation + "</Relationships>")
                data = text.encode("utf-8")
            elif info.filename == "[Content_Types].xml":
                text = data.decode("utf-8")
                if 'Extension="odttf"' not in text:
                    default = (
                        '<Default Extension="odttf" '
                        'ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont"/>'
                    )
                    text = text.replace("</Types>", default + "</Types>")
                data = text.encode("utf-8")
            zout.writestr(info, data)
        if relations_name not in source_names:
            relations = (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rIdNotoSansTC" '
                'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/font" '
                'Target="fonts/NotoSansTC.odttf"/></Relationships>'
            )
            zout.writestr(relations_name, relations.encode("utf-8"))
        zout.writestr("word/fonts/NotoSansTC.odttf", bytes(font_data))
    return target.getvalue()


def _display(value, percent: bool = False, signed: bool = False) -> str:
    if pd.isna(value):
        return "資料不足"
    if percent:
        number = float(value)
        return f"{number:+.2f}%" if signed else f"{number:.2f}%"
    if isinstance(value, (pd.Timestamp,)):
        return value.strftime("%Y-%m-%d")
    return str(value)


def _docx_set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    tc_pr.append(shading)


def _docx_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        _docx_set_cell_fill(cell, ACCENT)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
            for paragraph in cells[index].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
                    run.font.name = PDF_FONT
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), PDF_FONT)
    document.add_paragraph()


def build_docx_report(inputs, market_df, industry_df, alignment, holdings_df=None) -> bytes:
    holdings_df = holdings_df if holdings_df is not None else pd.DataFrame()
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = document.styles
    styles["Normal"].font.name = PDF_FONT
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), PDF_FONT)
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        styles[style_name].font.name = PDF_FONT
        styles[style_name]._element.rPr.rFonts.set(qn("w:eastAsia"), PDF_FONT)
        styles[style_name].font.color.rgb = RGBColor.from_string(ACCENT)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(f"{inputs['fund_name']} 一週基金分析報告")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        f"資料日期 {inputs['report_date']}  綜合判讀 {alignment}"
    ).bold = True
    document.add_paragraph(
        "本報告整合基金績效、正式 Benchmark、全球市場與基金前三大產業的近一週變化，"
        "用於快速核對基金表現與相關產業方向。"
    )

    document.add_heading("基金摘要", level=1)
    summary_rows = [
        ["基金本週", _display(inputs["fund_week"], True, True)],
        ["Benchmark 本週", _display(inputs["benchmark_week"], True, True)],
        ["超額報酬", _display(inputs["fund_week"] - inputs["benchmark_week"], True, True)],
        ["基金近一月", _display(inputs["fund_month"], True, True)],
        ["Sharpe", f"{inputs['sharpe']:.2f}"],
        ["Beta", f"{inputs['beta']:.2f}"],
        ["最大回撤", _display(inputs["max_drawdown"], True)],
        ["Benchmark", str(inputs["benchmark"])],
    ]
    _docx_table(document, ["指標", "數值"], summary_rows)

    document.add_heading("全球市場一週分析", level=1)
    from utils.live_data import market_summary

    document.add_paragraph(market_summary(market_df))
    market_rows = [
        [
            str(row.get("市場", "")),
            str(row.get("指數", "")),
            _display(row.get("本週漲跌%"), True, True),
            _display(row.get("近1月%"), True, True),
            str(row.get("趨勢", "")),
            _display(row.get("資料日期")),
        ]
        for _, row in market_df.iterrows()
    ]
    _docx_table(document, ["市場", "指數", "本週", "近一月", "趨勢", "日期"], market_rows)

    document.add_heading("持股前三大產業近一週變化", level=1)
    industry_rows = [
        [
            str(row.get("產業", "")),
            _display(row.get("持股權重%"), True),
            _display(row.get("本週變化%"), True, True),
            _display(row.get("近1月變化%"), True, True),
            str(row.get("趨勢", "")),
            _display(row.get("資料日期")),
        ]
        for _, row in industry_df.iterrows()
    ]
    _docx_table(document, ["產業", "權重", "本週", "近一月", "趨勢", "行情日期"], industry_rows)

    document.add_heading("前三大產業代表持股與公司", level=2)
    if holdings_df.empty:
        document.add_paragraph("逐檔行情資料不足。")
    else:
        holding_rows = [
            [
                str(row.get("產業", "")),
                str(row.get("公司／持股", "")),
                str(row.get("代碼", "")),
                _display(row.get("基金持股權重%"), True),
                _display(row.get("本週變化%"), True, True),
                _display(row.get("近1月變化%"), True, True),
                _display(row.get("行情日期")),
            ]
            for _, row in holdings_df.iterrows()
        ]
        _docx_table(
            document,
            ["產業", "公司", "代碼", "基金權重", "本週", "近一月", "日期"],
            holding_rows,
        )

    if inputs.get("news_content"):
        document.add_heading("新聞內容與市場觀察", level=1)
        document.add_paragraph(str(inputs["news_content"]))
        document.add_paragraph("本段為使用者提供內容，請核對原始來源與日期。")

    document.add_heading("資料來源與限制", level=1)
    document.add_paragraph(
        "全球市場與代表公司行情使用 Yahoo Finance 公開 chart 資料；基金名稱、Benchmark、"
        "淨值、風險指標與產業權重使用 MoneyDJ 公開資料。資料不足時不自行猜測。"
        "本報告僅供市場研究，不構成投資建議。"
    )
    buffer = BytesIO()
    document.save(buffer)
    return _embed_docx_font(buffer.getvalue())


def _pdf_table(headers: list[str], rows: list[list[str]], widths=None) -> Table:
    content = [[Paragraph(escape(item), ParagraphStyle("th", fontName=PDF_FONT, fontSize=8, textColor=colors.white)) for item in headers]]
    content.extend(
        [Paragraph(escape(str(item)), ParagraphStyle("td", fontName=PDF_FONT, fontSize=7.5, leading=10)) for item in row]
        for row in rows
    )
    table = Table(content, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{ACCENT}")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D8E0E8")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6FAF9")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return table


def build_pdf_report(inputs, market_df, industry_df, alignment, holdings_df=None) -> bytes:
    holdings_df = holdings_df if holdings_df is not None else pd.DataFrame()
    if PDF_FONT not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(PDF_FONT, str(PDF_FONT_PATH)))
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=0.48 * inch,
        leftMargin=0.48 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
        title=f"{inputs['fund_name']} 一週基金分析報告",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ChineseTitle",
        parent=styles["Title"],
        fontName=PDF_FONT,
        fontSize=21,
        leading=28,
        textColor=colors.HexColor(f"#{ACCENT}"),
        alignment=TA_CENTER,
        spaceAfter=12,
    )
    heading_style = ParagraphStyle(
        "ChineseHeading",
        fontName=PDF_FONT,
        fontSize=14,
        leading=18,
        textColor=colors.HexColor(f"#{ACCENT}"),
        spaceBefore=12,
        spaceAfter=8,
    )
    body_style = ParagraphStyle("ChineseBody", fontName=PDF_FONT, fontSize=9.5, leading=15)
    small_style = ParagraphStyle("ChineseSmall", fontName=PDF_FONT, fontSize=8, leading=12, textColor=colors.HexColor("#566274"))
    story = [
        Paragraph(escape(f"{inputs['fund_name']} 一週基金分析報告"), title_style),
        Paragraph(escape(f"資料日期 {inputs['report_date']}　綜合判讀 {alignment}"), body_style),
        Spacer(1, 8),
        Paragraph("基金摘要", heading_style),
    ]
    summary_rows = [
        ["基金本週", _display(inputs["fund_week"], True, True), "Benchmark 本週", _display(inputs["benchmark_week"], True, True)],
        ["超額報酬", _display(inputs["fund_week"] - inputs["benchmark_week"], True, True), "最大回撤", _display(inputs["max_drawdown"], True)],
        ["基金近一月", _display(inputs["fund_month"], True, True), "Sharpe／Beta", f"{inputs['sharpe']:.2f}／{inputs['beta']:.2f}"],
        ["Benchmark", str(inputs["benchmark"]), "綜合判讀", str(alignment)],
    ]
    story.append(_pdf_table(["指標", "數值", "指標", "數值"], summary_rows, [1.2*inch, 1.35*inch, 1.35*inch, 2.25*inch]))

    from utils.live_data import market_summary

    story.extend([Paragraph("全球市場一週分析", heading_style), Paragraph(escape(market_summary(market_df)), body_style)])
    market_rows = [[str(r.get("指數", "")), _display(r.get("本週漲跌%"), True, True), _display(r.get("近1月%"), True, True), str(r.get("趨勢", "")), _display(r.get("資料日期"))] for _, r in market_df.iterrows()]
    story.append(_pdf_table(["指數", "本週", "近一月", "趨勢", "日期"], market_rows, [2.0*inch, .8*inch, .8*inch, .8*inch, 1.1*inch]))

    story.append(Paragraph("持股前三大產業近一週變化", heading_style))
    industry_rows = [[str(r.get("產業", "")), _display(r.get("持股權重%"), True), _display(r.get("本週變化%"), True, True), _display(r.get("近1月變化%"), True, True), str(r.get("趨勢", "")), _display(r.get("資料日期"))] for _, r in industry_df.iterrows()]
    story.append(_pdf_table(["產業", "權重", "本週", "近一月", "趨勢", "日期"], industry_rows, [1.45*inch, .65*inch, .65*inch, .7*inch, .7*inch, 1.0*inch]))

    story.append(Paragraph("前三大產業代表持股與公司", heading_style))
    if holdings_df.empty:
        story.append(Paragraph("逐檔行情資料不足。", body_style))
    else:
        holding_rows = [[str(r.get("產業", "")), str(r.get("公司／持股", "")), str(r.get("代碼", "")), _display(r.get("基金持股權重%"), True), _display(r.get("本週變化%"), True, True), _display(r.get("近1月變化%"), True, True), _display(r.get("行情日期"))] for _, r in holdings_df.iterrows()]
        story.append(_pdf_table(["產業", "公司", "代碼", "權重", "本週", "近一月", "日期"], holding_rows, [1.05*inch, 1.0*inch, .85*inch, .6*inch, .6*inch, .65*inch, .9*inch]))

    if inputs.get("news_content"):
        story.extend([Paragraph("新聞內容與市場觀察", heading_style), Paragraph(escape(str(inputs["news_content"])).replace("\n", "<br/>"), body_style), Paragraph("本段為使用者提供內容，請核對原始來源與日期。", small_style)])
    story.extend([Paragraph("資料來源與限制", heading_style), Paragraph("全球市場與代表公司行情使用 Yahoo Finance 公開 chart 資料；基金資料使用 MoneyDJ 公開資料。資料不足不自行猜測。本報告僅供市場研究，不構成投資建議。", small_style)])
    doc.build(story)
    return buffer.getvalue()
